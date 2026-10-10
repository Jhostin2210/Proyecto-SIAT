"""
SIA-T: lectura del registro de notas del colegio (Excel).

El registro tiene una hoja por grado. En cada hoja hay bloques, uno por curso y mes:

    Nota:      | ALGEBRA 3ro |         | (fechas...)
    Noviembre  | APELLIDOS   | NOMBRES | LIBRO | CUADERNO | CONDUCTA | EXAMEN | PROMEDIO
    1          | ...         | ...     | 18    | 15       | 16       | 9      | 14.5

Este módulo recorre los bloques y devuelve las notas en formato largo, junto con los
problemas de calidad encontrados. No escribe en la base ni guarda nombres: el nombre
solo se usa como clave para asignar el código anónimo (ver anonimizar.py).
"""

import re
import unicodedata
from datetime import date, datetime

from openpyxl import load_workbook

MESES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "setiembre": 9, "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12,
}
NOMBRE_MES = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
ORDINALES = {"1": "1°", "2": "2°", "3": "3°", "4": "4°", "5": "5°", "6": "6°"}
MAX_COLUMNAS_NOTA = 6


def _texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def _plano(valor) -> str:
    """Minúsculas y sin tildes, para comparar encabezados."""
    t = unicodedata.normalize("NFD", _texto(valor).lower())
    return "".join(c for c in t if unicodedata.category(c) != "Mn")


def _grado_de_hoja(titulo: str):
    """'3ero PRIMARIA' -> ('3° Primaria', 'primaria', '3')."""
    plano = _plano(titulo)
    numero = re.search(r"(\d)", plano)
    nivel = "secundaria" if "sec" in plano else "primaria" if "prim" in plano else "inicial" if "inic" in plano else None
    if not numero or not nivel:
        return None
    n = numero.group(1)
    return f"{ORDINALES.get(n, n + '°')} {nivel.capitalize()}", nivel, n


def _componente(encabezado: str):
    plano = _plano(encabezado)
    for nombre in ("libro", "cuaderno", "conducta", "examen"):
        if plano.startswith(nombre):
            return nombre
    return None


# Nombre oficial de cada curso según cómo aparece escrito en el registro (sin tildes ni mayúsculas).
CURSOS = {
    "algebra": "Álgebra",
    "geometria": "Geometría",
    "aritmetica": "Aritmética",
    "aritmatica": "Aritmética",  # error de escritura frecuente en el registro
    "historia": "Historia",
    "economia": "Economía",
    "r.m": "Razonamiento matemático",
    "r.v": "Razonamiento verbal",
}


def _curso(titulo: str) -> str:
    """'ALGEBRA 3ro (E.F.)' -> 'Álgebra'. Usa el texto anterior al primer dígito o paréntesis."""
    base = re.split(r"[\d(]", _texto(titulo))[0].strip(" .-")
    return CURSOS.get(_plano(base), base[:1].upper() + base[1:].lower())


def leer_registro(ruta: str) -> dict:
    """Devuelve {"notas": [...], "hallazgos": [...], "filas_origen": n, "bloques": n, "resumen": {...}}."""
    libro = load_workbook(ruta, data_only=True)
    notas, hallazgos, vistos = [], [], set()
    filas_origen = bloques = 0
    estructura = {}

    def hallazgo(regla, severidad, detalle):
        hallazgos.append({"regla": regla, "severidad": severidad, "detalle": detalle})

    for hoja in libro.worksheets:
        info = _grado_de_hoja(hoja.title)
        if not info:
            hallazgo("hoja_no_reconocida", "aviso", f"La hoja «{hoja.title}» no indica un grado y no se leyó.")
            continue
        grado, nivel, numero_grado = info
        celdas = [[c for c in fila] for fila in hoja.iter_rows(values_only=True)]
        alto = len(celdas)

        def celda(f, c):
            return celdas[f][c] if 0 <= f < alto and 0 <= c < len(celdas[f]) else None

        for f in range(alto):
            for c in range(len(celdas[f]) - 1):
                if _plano(celda(f, c)) != "apellidos" or _plano(celda(f, c + 1)) != "nombres":
                    continue
                bloques += 1
                ubicacion = f"hoja «{hoja.title}», fila {f + 1}"
                titulo = _texto(celda(f - 1, c))
                mes = MESES.get(_plano(celda(f, c - 1)))
                anio = next((v.year for v in celdas[f - 1][c:c + 8] if isinstance(v, (date, datetime))), None) if f > 0 else None
                if not mes or not anio:
                    hallazgo("periodo_no_identificado", "critica", f"Bloque sin mes o año reconocible ({ubicacion}); no se cargó.")
                    continue
                periodo = f"{anio}-{mes:02d}"
                curso = _curso(titulo)
                if _plano(re.split(r"[\d(]", titulo)[0].strip(" .-")) not in CURSOS:
                    hallazgo("curso_no_catalogado", "aviso", f"El curso «{curso}» no está en el catálogo ({ubicacion}); se cargó con ese nombre.")
                numero_titulo = re.search(r"(\d)", titulo)
                if numero_titulo and numero_titulo.group(1) != numero_grado:
                    hallazgo("grado_del_titulo_distinto", "aviso", f"El título «{titulo}» indica otro grado que la hoja ({ubicacion}). Se usó el grado de la hoja.")

                # Columnas de notas: hasta el primer encabezado que no sea un componente (el promedio).
                columnas = []
                for k in range(c + 2, c + 2 + MAX_COLUMNAS_NOTA):
                    encabezado = _texto(celda(f, k))
                    if _plano(encabezado) in ("apellidos", "nota:", "nota."):
                        break
                    comp = _componente(encabezado)
                    if comp and comp not in [x for _, x in columnas]:
                        columnas.append((k, comp))
                    elif comp:
                        hallazgo("componente_repetido", "aviso", f"«{encabezado}» repite un componente y no se cargó ({ubicacion}).")
                    elif not encabezado:
                        if any(isinstance(celda(r, k), (int, float)) for r in range(f + 1, min(f + 4, alto))):
                            hallazgo("columna_sin_encabezado", "aviso", f"Hay una columna con notas pero sin encabezado ({ubicacion}); no se cargó.")
                    else:
                        break

                clave_bloque = (grado, curso, periodo)
                if clave_bloque in vistos:
                    hallazgo("bloque_repetido", "aviso", f"{curso}, {grado}, {periodo} aparece más de una vez ({ubicacion}). Se conservó el primero.")
                    continue
                vistos.add(clave_bloque)
                estructura.setdefault(grado, {"nivel": nivel, "cursos": {}, "estudiantes": set()})
                estructura[grado]["cursos"].setdefault(curso, {"componentes": [x for _, x in columnas], "periodos": []})["periodos"].append(periodo)

                r = f + 1
                while r < alto and _texto(celda(r, c)):
                    filas_origen += 1
                    clave = f"{_plano(celda(r, c))}|{_plano(celda(r, c + 1))}|{grado}"
                    valores = [(comp, celda(r, k)) for k, comp in columnas]
                    numericos = [v for _, v in valores if isinstance(v, (int, float))]
                    if numericos and all(v == 0 for v in numericos):
                        hallazgo("fila_en_cero", "aviso", f"Estudiante con todas las notas en cero en {curso}, {grado}, {periodo} (fila {r + 1}); no se cargó.")
                        r += 1
                        continue
                    estructura[grado]["estudiantes"].add(clave)
                    for comp, v in valores:
                        if v is None or _texto(v) == "":
                            hallazgo("nota_vacia", "aviso", f"Nota de {comp} vacía en {curso}, {grado}, {periodo} (fila {r + 1}).")
                        elif not isinstance(v, (int, float)):
                            hallazgo("nota_no_numerica", "critica", f"Valor no numérico en {comp}, {curso}, {grado}, {periodo} (fila {r + 1}); no se cargó.")
                        elif not 0 <= v <= 20:
                            hallazgo("nota_fuera_de_rango", "critica", f"Nota {v} fuera de 0 a 20 en {comp}, {curso}, {grado}, {periodo} (fila {r + 1}); no se cargó.")
                        else:
                            if v == 0:
                                hallazgo("nota_en_cero", "aviso", f"Nota 0 en {comp}, {curso}, {grado}, {periodo} (fila {r + 1}); se cargó, conviene confirmarla.")
                            notas.append({"clave": clave, "grado": grado, "nivel": nivel, "curso": curso, "periodo": periodo, "componente": comp, "nota": float(v)})
                    r += 1

    resumen = {
        g: {
            "nivel": d["nivel"],
            "estudiantes": len(d["estudiantes"]),
            "cursos": {cu: {"componentes": v["componentes"], "periodos": sorted(set(v["periodos"]))} for cu, v in d["cursos"].items()},
        }
        for g, d in estructura.items()
    }
    return {"notas": notas, "hallazgos": hallazgos, "filas_origen": filas_origen, "bloques": bloques, "resumen": resumen}


def nombre_periodo(codigo: str) -> str:
    anio, mes = codigo.split("-")
    return f"{NOMBRE_MES[int(mes)]} de {anio}"
